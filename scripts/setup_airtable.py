"""
Setup automático de la base Airtable para el PoC de Career Development.
Crea: base, 5 tablas con campos y relaciones, y migra el catalogo de skills
desde el Excel original.

Requisitos:
    pip install requests openpyxl

Uso:
    set AIRTABLE_TOKEN=pat...                 (PowerShell: $env:AIRTABLE_TOKEN="pat...")
    set AIRTABLE_WORKSPACE_ID=wsp...          (PowerShell: $env:AIRTABLE_WORKSPACE_ID="wsp...")
    python setup_airtable.py
"""

import os
import time
import requests
from openpyxl import load_workbook

TOKEN = os.environ["AIRTABLE_TOKEN"]
WORKSPACE_ID = os.environ["AIRTABLE_WORKSPACE_ID"]
EXCEL_PATH = r"C:\Users\christian.rojas\Downloads\Career Development Tool COPY.xlsx"
BASE_NAME = "Career Development PoC"

API_ROOT = "https://api.airtable.com/v0"
HEADERS = {"Authorization": f"Bearer {TOKEN}", "Content-Type": "application/json"}


def create_base():
    body = {
        "name": BASE_NAME,
        "workspaceId": WORKSPACE_ID,
        "tables": [
            {
                "name": "Catalogo_Skills",
                "fields": [
                    {"name": "Nombre_Skill", "type": "singleLineText"},
                    {"name": "Tipo", "type": "singleSelect", "options": {"choices": [
                        {"name": "Funcion"}, {"name": "Soft Skill"}, {"name": "Hard Skill"}]}},
                    {"name": "Perfil_Rol", "type": "singleLineText"},
                    {"name": "Descripcion", "type": "multilineText"},
                    {"name": "Nivel_Basic", "type": "multilineText"},
                    {"name": "Nivel_Intermediate", "type": "multilineText"},
                    {"name": "Nivel_Advanced", "type": "multilineText"},
                    {"name": "Nivel_Expert", "type": "multilineText"},
                ],
            },
            {
                "name": "Personas",
                "fields": [
                    {"name": "Nombre", "type": "singleLineText"},
                    {"name": "Area", "type": "singleLineText"},
                    {"name": "Perfil_Rol", "type": "singleLineText"},
                    {"name": "Email", "type": "email"},
                ],
            },
            {
                "name": "Rondas",
                "fields": [
                    {"name": "Nombre_Ronda", "type": "singleLineText"},
                    {"name": "Fecha_Inicio", "type": "date", "options": {"dateFormat": {"name": "iso"}}},
                    {"name": "Fecha_Cierre", "type": "date", "options": {"dateFormat": {"name": "iso"}}},
                    {"name": "Estado", "type": "singleSelect", "options": {"choices": [
                        {"name": "Abierta"}, {"name": "Cerrada"}]}},
                ],
            },
        ],
    }
    r = requests.post(f"{API_ROOT}/meta/bases", headers=HEADERS, json=body)
    if r.status_code >= 300:
        print("ERROR creando base:", r.status_code, r.text)
    r.raise_for_status()
    data = r.json()
    base_id = data["id"]
    tables = {t["name"]: t["id"] for t in data["tables"]}
    print("Base creada:", base_id)
    return base_id, tables


def add_link_field(base_id, table_id, field_name, target_table_id, extra=None):
    body = {"name": field_name, "type": "multipleRecordLinks",
            "options": {"linkedTableId": target_table_id}}
    if extra:
        body.update(extra)
    r = requests.post(f"{API_ROOT}/meta/bases/{base_id}/tables/{table_id}/fields",
                       headers=HEADERS, json=body)
    r.raise_for_status()
    return r.json()


def add_field(base_id, table_id, body):
    r = requests.post(f"{API_ROOT}/meta/bases/{base_id}/tables/{table_id}/fields",
                       headers=HEADERS, json=body)
    r.raise_for_status()
    return r.json()


def create_evaluaciones_table(base_id, tables):
    body = {
        "name": "Evaluaciones",
        "fields": [{"name": "Nombre_Evaluador", "type": "singleLineText"}],
    }
    r = requests.post(f"{API_ROOT}/meta/bases/{base_id}/tables", headers=HEADERS, json=body)
    r.raise_for_status()
    table = r.json()
    table_id = table["id"]
    tables["Evaluaciones"] = table_id

    add_link_field(base_id, table_id, "Ronda", tables["Rondas"])
    add_link_field(base_id, table_id, "Evaluado", tables["Personas"])
    add_link_field(base_id, table_id, "Skill", tables["Catalogo_Skills"])
    add_field(base_id, table_id, {"name": "Rol_Evaluador", "type": "singleSelect",
              "options": {"choices": [{"name": "Employee"}, {"name": "Leader"}, {"name": "Peer"}]}})
    add_field(base_id, table_id, {"name": "Puntaje", "type": "number",
              "options": {"precision": 0}})
    print("Tabla Evaluaciones creada:", table_id)


def create_plan_desarrollo_table(base_id, tables):
    body = {
        "name": "Plan_Desarrollo",
        "fields": [{"name": "Accion_Recomendada", "type": "multilineText"}],
    }
    r = requests.post(f"{API_ROOT}/meta/bases/{base_id}/tables", headers=HEADERS, json=body)
    r.raise_for_status()
    table = r.json()
    table_id = table["id"]
    tables["Plan_Desarrollo"] = table_id

    add_link_field(base_id, table_id, "Persona", tables["Personas"])
    add_link_field(base_id, table_id, "Skill_Gap", tables["Catalogo_Skills"])
    add_field(base_id, table_id, {"name": "Responsable", "type": "singleLineText"})
    add_field(base_id, table_id, {"name": "Fecha_Compromiso", "type": "date",
              "options": {"dateFormat": {"name": "iso"}}})
    add_field(base_id, table_id, {"name": "Estado", "type": "singleSelect", "options": {"choices": [
        {"name": "No iniciado"}, {"name": "En curso"}, {"name": "Completado"}]}})
    print("Tabla Plan_Desarrollo creada:", table_id)


def add_lider_selflink(base_id, tables):
    add_link_field(base_id, tables["Personas"], "Lider", tables["Personas"])
    print("Campo Lider (auto-referencia) agregado a Personas")


def read_catalog_from_excel():
    """
    Columnas reales verificadas en el Excel origen:
      Functions:   B=Key C=Department D=Function     E=Description F..I=Basic..Expert (fila 7 en adelante)
      Hard Skills: A=Column1 B=Department C=Competency D=Description  E..H=Basic..Expert (fila 4 en adelante)
      Soft Skills: C=Soft Skill D=Description E..H=Basic..Expert (fila 6 en adelante)
    """
    wb = load_workbook(EXCEL_PATH, data_only=True)
    rows = []

    ws = wb["Functions"]
    for row in ws.iter_rows(min_row=7, max_row=115, min_col=2, max_col=9):
        vals = [c.value for c in row]  # B,C,D,E,F,G,H,I -> idx 0..7
        if not vals[2]:
            continue
        rows.append({
            "Nombre_Skill": vals[2], "Tipo": "Funcion", "Perfil_Rol": vals[1] or "",
            "Descripcion": vals[3] or "", "Nivel_Basic": vals[4] or "",
            "Nivel_Intermediate": vals[5] or "", "Nivel_Advanced": vals[6] or "",
            "Nivel_Expert": vals[7] or "",
        })

    ws = wb["Hard Skills"]
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=2, max_col=8):
        vals = [c.value for c in row]  # B,C,D,E,F,G,H -> idx 0..6
        if not vals[1]:
            continue
        rows.append({
            "Nombre_Skill": vals[1], "Tipo": "Hard Skill", "Perfil_Rol": vals[0] or "",
            "Descripcion": vals[2] or "", "Nivel_Basic": vals[3] or "",
            "Nivel_Intermediate": vals[4] or "", "Nivel_Advanced": vals[5] or "",
            "Nivel_Expert": vals[6] or "",
        })

    ws = wb["Soft Skills"]
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, min_col=3, max_col=8):
        vals = [c.value for c in row]  # C,D,E,F,G,H -> idx 0..5
        if not vals[0]:
            continue
        rows.append({
            "Nombre_Skill": vals[0], "Tipo": "Soft Skill", "Perfil_Rol": "",
            "Descripcion": vals[1] or "", "Nivel_Basic": vals[2] or "",
            "Nivel_Intermediate": vals[3] or "", "Nivel_Advanced": vals[4] or "",
            "Nivel_Expert": vals[5] or "",
        })

    return rows


def push_records(base_id, table_id, records, table_name):
    url = f"{API_ROOT}/{base_id}/{table_name}"
    for i in range(0, len(records), 10):
        batch = records[i:i + 10]
        body = {"records": [{"fields": r} for r in batch]}
        r = requests.post(url, headers=HEADERS, json=body)
        if r.status_code >= 300:
            print("ERROR:", r.status_code, r.text)
        r.raise_for_status()
        time.sleep(0.25)
    print(f"{len(records)} registros cargados en {table_name}")


if __name__ == "__main__":
    base_id, tables = create_base()
    create_evaluaciones_table(base_id, tables)
    create_plan_desarrollo_table(base_id, tables)
    add_lider_selflink(base_id, tables)

    print("\nIDs de tablas:", tables)
    print("BASE_ID:", base_id, "\n-> Guarda estos valores, los necesitas para los HTML.\n")

    catalog_rows = read_catalog_from_excel()
    push_records(base_id, tables["Catalogo_Skills"], catalog_rows, "Catalogo_Skills")

    print("\nListo. Base:", f"https://airtable.com/{base_id}")
